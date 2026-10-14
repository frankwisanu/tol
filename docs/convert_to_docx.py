"""
Convert Markdown documentation to professional Microsoft Word (.docx) & WPS Office format.
Project: True Online (TOL) Automation Testing Standard & Role Deliverables
"""

import os
import sys
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

PRIMARY_NAVY = "0F4C81"      # Deep Telecom / Enterprise Navy
SECONDARY_BLUE = "1A5276"    # Medium Blue
ACCENT_BLUE = "2E86C1"       # Light Accent Blue
BG_LIGHT = "F4F6F7"          # Light Gray / Off-White for cards & code
ZEBRA_ROW = "F8F9FA"         # Zebra striping for tables
BORDER_GRAY = "D5D8DC"       # Subtle table borders
CODE_RED = RGBColor(199, 37, 78)  # Monospace inline code color
TEXT_DARK = RGBColor(40, 40, 40)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>\n'
        f'  <w:top w:w="{top}" w:type="dxa"/>\n'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>\n'
        f'  <w:left w:w="{left}" w:type="dxa"/>\n'
        f'  <w:right w:w="{right}" w:type="dxa"/>\n'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top=BORDER_GRAY, bottom=BORDER_GRAY, left=BORDER_GRAY, right=BORDER_GRAY, sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{top}"/>\n'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{bottom}"/>\n'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{left}"/>\n'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{right}"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def set_left_accent_border(cell, color=PRIMARY_NAVY, sz="24"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def add_formatted_text(paragraph, text, base_font="Calibri", base_size=Pt(10.5), text_color=TEXT_DARK, is_italic=False):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.name = base_font
            run.font.size = base_size
            run.font.color.rgb = text_color
            if is_italic:
                run.italic = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.font.name = base_font
            run.font.size = base_size
            run.font.color.rgb = text_color
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = CODE_RED
        else:
            run = paragraph.add_run(token)
            run.font.name = base_font
            run.font.size = base_size
            run.font.color.rgb = text_color
            if is_italic:
                run.italic = True

def parse_markdown_to_docx(md_path, docx_path, doc_title="TrueOnline Specification"):
    print(f"Converting: {os.path.basename(md_path)} -> {os.path.basename(docx_path)}")
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # Page Margins: 1 inch all sides
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

        # Header & Footer setup
        header = section.header
        hp = header.paragraphs[0]
        hp.text = f"True Online (TOL) Automation Testing Standard  |  {doc_title}"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.runs[0].font.name = "Calibri"
        hp.runs[0].font.size = Pt(8.5)
        hp.runs[0].font.color.rgb = RGBColor(130, 140, 150)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Confidential  •  For Stakeholder Review, Visual Studio Presentation & Testing Teams"
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp.runs[0].font.name = "Calibri"
        fp.runs[0].font.size = Pt(8.5)
        fp.runs[0].font.color.rgb = RGBColor(130, 140, 150)

    # Document Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = TEXT_DARK
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    def flush_code_block():
        nonlocal code_lines
        if not code_lines:
            return
        code_text = "".join(code_lines)
        code_lines = []
        # Create 1x1 table for shaded box
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, BG_LIGHT)
        set_cell_borders(cell, top="CCD1D9", bottom="CCD1D9", left="CCD1D9", right="CCD1D9", sz="4")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        run = p.add_run(code_text.rstrip("\n"))
        run.font.name = "Consolas"
        run.font.size = Pt(9.0)
        run.font.color.rgb = RGBColor(35, 45, 55)

        # Spacing after table
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(4)

    def flush_table():
        nonlocal table_rows
        if not table_rows:
            return
        
        # Parse table rows into cells
        parsed_rows = []
        for r_line in table_rows:
            raw_cells = [c.strip() for c in r_line.strip().strip('|').split('|')]
            parsed_rows.append(raw_cells)
        table_rows = []

        if not parsed_rows:
            return

        header_cells = parsed_rows[0]
        data_cells = parsed_rows[1:]
        col_count = len(header_cells)

        tbl = doc.add_table(rows=len(parsed_rows), cols=col_count)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        # Calculate approximate column widths
        col_widths = [Inches(6.5 / col_count)] * col_count
        # If 2 columns, make col 1 2.2 in, col 2 4.3 in
        if col_count == 2:
            col_widths = [Inches(2.2), Inches(4.3)]
        elif col_count == 3:
            col_widths = [Inches(1.8), Inches(1.8), Inches(2.9)]
        elif col_count == 4:
            col_widths = [Inches(1.5), Inches(1.5), Inches(1.8), Inches(1.7)]
        elif col_count == 5:
            col_widths = [Inches(1.1), Inches(1.3), Inches(1.6), Inches(1.2), Inches(1.3)]
        elif col_count == 7:
            col_widths = [Inches(0.6), Inches(1.3), Inches(0.9), Inches(0.7), Inches(1.1), Inches(1.0), Inches(0.9)]

        # Header Row
        hdr_row = tbl.rows[0]
        make_row_header(hdr_row)
        for c_idx, cell_text in enumerate(header_cells):
            cell = hdr_row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, PRIMARY_NAVY)
            set_cell_borders(cell, top=PRIMARY_NAVY, bottom=PRIMARY_NAVY, left="205C91", right="205C91", sz="4")
            set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_formatted_text(p, cell_text, base_font="Calibri", base_size=Pt(10.0), text_color=RGBColor(255, 255, 255))
            for r in p.runs:
                r.bold = True

        # Data Rows
        for r_idx, row_data in enumerate(data_cells):
            row = tbl.rows[r_idx + 1]
            prevent_row_split(row)
            bg_color = ZEBRA_ROW if (r_idx % 2 == 1) else "FFFFFF"
            for c_idx in range(col_count):
                cell_text = row_data[c_idx] if c_idx < len(row_data) else ""
                cell = row.cells[c_idx]
                cell.width = col_widths[c_idx]
                set_cell_background(cell, bg_color)
                set_cell_borders(cell, top=BORDER_GRAY, bottom=BORDER_GRAY, left=BORDER_GRAY, right=BORDER_GRAY, sz="4")
                set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                add_formatted_text(p, cell_text, base_font="Calibri", base_size=Pt(9.5))

        # Add small spacing after table
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(6)

    # Line by line processing
    idx = 0
    total_lines = len(lines)
    while idx < total_lines:
        line = lines[idx]
        stripped = line.strip()

        # Handle Code Fences ```
        if stripped.startswith("```"):
            if in_table:
                flush_table()
                in_table = False
            if in_code_block:
                in_code_block = False
                flush_code_block()
            else:
                in_code_block = True
                code_lines = []
            idx += 1
            continue

        if in_code_block:
            code_lines.append(line)
            idx += 1
            continue

        # Handle Markdown Tables (| ... |)
        if stripped.startswith("|") and stripped.endswith("|"):
            # Check if separator row (|---|---|)
            if re.match(r'^\|[\s\-:|]+\|$', stripped):
                idx += 1
                continue
            in_table = True
            table_rows.append(stripped)
            idx += 1
            continue
        else:
            if in_table:
                flush_table()
                in_table = False

        # Blank line
        if not stripped:
            idx += 1
            continue

        # Horizontal Rule (---)
        if stripped in ["---", "***", "___"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="{PRIMARY_NAVY}"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            idx += 1
            continue

        # Headings #, ##, ###, ####
        if stripped.startswith("#"):
            match = re.match(r'^(#{1,4})\s+(.*)', stripped)
            if match:
                level = len(match.group(1))
                h_text = match.group(2)
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                if level == 1:
                    p.paragraph_format.space_before = Pt(16)
                    p.paragraph_format.space_after = Pt(6)
                    add_formatted_text(p, h_text, base_font="Calibri", base_size=Pt(18.0), text_color=RGBColor(15, 76, 129))
                    for r in p.runs:
                        r.bold = True
                    # subtle bottom underline for H1
                    p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="{PRIMARY_NAVY}"/></w:pBdr>')
                    p._p.get_or_add_pPr().append(p_border)
                elif level == 2:
                    p.paragraph_format.space_before = Pt(13)
                    p.paragraph_format.space_after = Pt(4)
                    add_formatted_text(p, h_text, base_font="Calibri", base_size=Pt(14.0), text_color=RGBColor(26, 82, 118))
                    for r in p.runs:
                        r.bold = True
                elif level == 3:
                    p.paragraph_format.space_before = Pt(9)
                    p.paragraph_format.space_after = Pt(3)
                    add_formatted_text(p, h_text, base_font="Calibri", base_size=Pt(12.0), text_color=RGBColor(46, 134, 193))
                    for r in p.runs:
                        r.bold = True
                else:
                    p.paragraph_format.space_before = Pt(7)
                    p.paragraph_format.space_after = Pt(2)
                    add_formatted_text(p, h_text, base_font="Calibri", base_size=Pt(11.0), text_color=RGBColor(52, 73, 94))
                    for r in p.runs:
                        r.bold = True
                idx += 1
                continue

        # Blockquote (> ...)
        if stripped.startswith(">"):
            bq_text = stripped.lstrip(">").strip()
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False
            cell = tbl.cell(0, 0)
            cell.width = Inches(6.5)
            set_cell_background(cell, BG_LIGHT)
            set_left_accent_border(cell, color=PRIMARY_NAVY, sz="24")
            set_cell_margins(cell, top=100, bottom=100, left=160, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            add_formatted_text(p, bq_text, base_font="Calibri", base_size=Pt(10.5), text_color=RGBColor(50, 60, 70), is_italic=True)
            sp = doc.add_paragraph()
            sp.paragraph_format.space_before = Pt(0)
            sp.paragraph_format.space_after = Pt(4)
            idx += 1
            continue

        # Bullet List (- or *)
        bullet_match = re.match(r'^(\s*)[-*]\s+(.*)', line)
        if bullet_match:
            indent_spaces = len(bullet_match.group(1))
            bullet_text = bullet_match.group(2)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if indent_spaces >= 2:
                p.paragraph_format.left_indent = Inches(0.5)
            add_formatted_text(p, bullet_text, base_font="Calibri", base_size=Pt(10.0))
            idx += 1
            continue

        # Numbered List (1. ...)
        num_match = re.match(r'^(\s*)(\d+)\.\s+(.*)', line)
        if num_match:
            indent_spaces = len(num_match.group(1))
            num_str = num_match.group(2)
            num_text = num_match.group(3)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            # Add prefix number
            run_num = p.add_run(f"{num_str}.  ")
            run_num.bold = True
            run_num.font.name = "Calibri"
            run_num.font.size = Pt(10.0)
            run_num.font.color.rgb = RGBColor(15, 76, 129)
            add_formatted_text(p, num_text, base_font="Calibri", base_size=Pt(10.0))
            idx += 1
            continue

        # Regular Body Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_text(p, stripped, base_font="Calibri", base_size=Pt(10.5))
        idx += 1

    # End of file flushes
    if in_table:
        flush_table()
    if in_code_block:
        flush_code_block()

    doc.save(docx_path)
    print(f"Successfully generated: {docx_path}")

def main():
    docs_dir = os.path.dirname(os.path.abspath(__file__))
    root_dir = os.path.dirname(docs_dir)

    targets = [
        (os.path.join(docs_dir, "01_BA_Business_Analysis_Specification.md"),
         os.path.join(docs_dir, "01_BA_Business_Analysis_Specification.docx"),
         "BA Business Analysis Specification"),

        (os.path.join(docs_dir, "02_PO_Product_Owner_Specification.md"),
         os.path.join(docs_dir, "02_PO_Product_Owner_Specification.docx"),
         "PO Product Owner Specification"),

        (os.path.join(docs_dir, "03_TOL_Requirements_Document_BRD_PRD.md"),
         os.path.join(docs_dir, "03_TOL_Requirements_Document_BRD_PRD.docx"),
         "TOL Requirements Document (BRD / PRD)"),

        (os.path.join(docs_dir, "04_Traceability_Matrix_And_Test_Mapping.md"),
         os.path.join(docs_dir, "04_Traceability_Matrix_And_Test_Mapping.docx"),
         "E2E Traceability Matrix & Test Mapping"),

        (os.path.join(docs_dir, "05_Deployment_And_Hosting_Guide_FrankTest.md"),
         os.path.join(docs_dir, "05_Deployment_And_Hosting_Guide_FrankTest.docx"),
         "Deployment & Hosting Guide"),

        (os.path.join(root_dir, "README.md"),
         os.path.join(docs_dir, "00_TOL_Automation_Testing_Standard_Overview.docx"),
         "TOL Automation Testing Standard & Architecture Overview"),

        (os.path.join(docs_dir, "07_Sprint_Plan_Week_12Oct2026.md"),
         os.path.join(docs_dir, "07_Sprint_Plan_Week_12Oct2026.docx"),
         "Sprint Plan - Week of 12 October 2026"),
    ]

    for md_file, docx_file, title in targets:
        if os.path.exists(md_file):
            parse_markdown_to_docx(md_file, docx_file, title)
        else:
            print(f"File not found: {md_file}")

if __name__ == "__main__":
    main()
